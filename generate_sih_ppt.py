import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_sih_slide(output_path):
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)
    
    # Colors
    c_white = RGBColor(255, 255, 255)
    c_primary = RGBColor(30, 27, 75)         # Indigo 950
    c_accent_blue = RGBColor(37, 99, 235)    # Blue 600
    c_border = RGBColor(203, 213, 225)       # Slate 300
    c_text_dark = RGBColor(15, 23, 42)
    c_text_muted = RGBColor(71, 85, 105)
    c_green = RGBColor(22, 101, 52)          # Green
    c_red = RGBColor(185, 28, 28)            # Red
    c_link = RGBColor(29, 78, 216)           # Blue Link
    
    # Slide Background
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bg_shape.line.fill.background()
    
    # Top Header Bar / SIH Tag
    sih_badge = slide.shapes.add_textbox(Inches(11.2), Inches(0.2), Inches(1.8), Inches(0.5))
    sih_tf = sih_badge.text_frame
    sih_tf.word_wrap = True
    p_sih = sih_tf.paragraphs[0]
    p_sih.text = "SIH 2026"
    p_sih.font.bold = True
    p_sih.font.size = Pt(16)
    p_sih.font.color.rgb = c_primary
    p_sih.alignment = PP_ALIGN.RIGHT

    # Project Title / Subtitle
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.15), Inches(10), Inches(0.55))
    t_tf = title_box.text_frame
    t_p = t_tf.paragraphs[0]
    t_p.text = "SatQuery AI  |  Bi-Temporal Satellite Intelligence & VQA"
    t_p.font.bold = True
    t_p.font.size = Pt(15)
    t_p.font.color.rgb = c_accent_blue

    # =========================================================================
    # LEFT COLUMN: REFERENCES & WORKFLOW
    # =========================================================================
    
    # Left Card Container
    left_card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.75), Inches(5.8), Inches(6.3)
    )
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = RGBColor(248, 250, 252)
    left_card.line.color.rgb = c_border
    left_card.line.width = Pt(1)
    
    # Left Header: ❖ References
    ref_head = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(5.4), Inches(0.45))
    rf_tf = ref_head.text_frame
    rf_p = rf_tf.paragraphs[0]
    rf_p.text = "❖ References & Research Foundations"
    rf_p.font.bold = True
    rf_p.font.size = Pt(15)
    rf_p.font.color.rgb = c_primary
    
    # References Content Box
    ref_body = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(5.4), Inches(3.3))
    rb_tf = ref_body.text_frame
    rb_tf.word_wrap = True
    
    ref_sections = [
        ("Satellite Data Platforms:", [
            ("➢ Copernicus Data Space (Sentinel-1 SAR / Sentinel-2):", "https://dataspace.copernicus.eu/"),
            ("➢ USGS EarthExplorer (Landsat-8/9 Multispectral):", "https://earthexplorer.usgs.gov/"),
            ("➢ ESA Sentinel Hub Technical Guides:", "https://sentinels.copernicus.eu/")
        ]),
        ("Research & Benchmark Datasets:", [
            ("➢ LEVIR-CD (Bi-Temporal Change Detection):", "https://chenhao.in/LEVIR/"),
            ("➢ BigEarthNet (Multispectral Remote Sensing):", "https://bigearth.net/"),
            ("➢ RSVQA: Visual Question Answering (IEEE TGRS):", "https://rsvqa.sylvainlobry.com/")
        ]),
        ("Feasibility & Geospatial Standards:", [
            ("➢ OGC GeoTIFF & CRS Coordinate Standards:", "https://www.ogc.org/standards/geotiff"),
            ("➢ ISRO Bhuvan Open Geoportal:", "https://bhuvan.nrsc.gov.in/")
        ])
    ]
    
    first = True
    for header, items in ref_sections:
        p_hdr = rb_tf.paragraphs[0] if first else rb_tf.add_paragraph()
        first = False
        p_hdr.text = header
        p_hdr.font.bold = True
        p_hdr.font.size = Pt(10)
        p_hdr.font.color.rgb = RGBColor(17, 24, 39)
        p_hdr.space_before = Pt(4)
        p_hdr.space_after = Pt(2)
        
        for title, url in items:
            p_item = rb_tf.add_paragraph()
            p_item.text = f"{title} {url}"
            p_item.font.size = Pt(8.5)
            p_item.font.color.rgb = c_link
            p_item.space_after = Pt(1)

    # Workflow Section Header
    wf_head = slide.shapes.add_textbox(Inches(0.8), Inches(4.7), Inches(5.4), Inches(0.35))
    wf_tf = wf_head.text_frame
    wf_p = wf_tf.paragraphs[0]
    wf_p.text = "❖ Research & System Workflow"
    wf_p.font.bold = True
    wf_p.font.size = Pt(12)
    wf_p.font.color.rgb = c_primary

    # Flowchart Visual Steps
    steps = [
        ("T1 & T2 Ingestion", "Rasterio & CRS validation"),
        ("Siamese Change Net", "Bounding box & mask extraction"),
        ("Spectral Transition", "NDVI, NDWI, NDBI deltas"),
        ("Evidence & SAR Audit", "Zero-hallucination trace")
    ]
    
    step_width = Inches(1.22)
    step_height = Inches(1.3)
    start_x = Inches(0.8)
    y_pos = Inches(5.1)
    
    for i, (st_title, st_desc) in enumerate(steps):
        s_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, start_x + i * Inches(1.33), y_pos, step_width, step_height
        )
        s_box.fill.solid()
        s_box.fill.fore_color.rgb = RGBColor(238, 242, 255) if i % 2 == 0 else RGBColor(240, 253, 244)
        s_box.line.color.rgb = c_accent_blue if i % 2 == 0 else RGBColor(34, 197, 94)
        s_box.line.width = Pt(1)
        
        s_tf = s_box.text_frame
        s_tf.word_wrap = True
        s_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        
        p1 = s_tf.paragraphs[0]
        p1.text = f"Step {i+1}\n{st_title}"
        p1.font.bold = True
        p1.font.size = Pt(8.5)
        p1.font.color.rgb = c_primary
        p1.alignment = PP_ALIGN.CENTER
        
        p2 = s_tf.add_paragraph()
        p2.text = st_desc
        p2.font.size = Pt(7.5)
        p2.font.color.rgb = c_text_muted
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(2)

    # =========================================================================
    # RIGHT COLUMN: COMPARISON WITH EXISTING SYSTEMS & LIVE DEMO
    # =========================================================================
    
    # Right Card Container
    right_card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.65), Inches(0.75), Inches(6.1), Inches(6.3)
    )
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = RGBColor(255, 255, 255)
    right_card.line.color.rgb = c_border
    right_card.line.width = Pt(1)

    # Right Header: ❖ Comparison with Existing Systems
    comp_head = slide.shapes.add_textbox(Inches(6.8), Inches(0.85), Inches(5.8), Inches(0.45))
    cf_tf = comp_head.text_frame
    cf_p = cf_tf.paragraphs[0]
    cf_p.text = "❖ Comparison with Existing Systems"
    cf_p.font.bold = True
    cf_p.font.size = Pt(15)
    cf_p.font.color.rgb = c_primary
    
    # Comparison Table
    rows = 9
    cols = 5
    table_shape = slide.shapes.add_table(
        rows, cols, Inches(6.8), Inches(1.35), Inches(5.8), Inches(4.5)
    )
    tbl = table_shape.table
    
    # Column Widths
    tbl.columns[0].width = Inches(2.3)
    tbl.columns[1].width = Inches(0.9)
    tbl.columns[2].width = Inches(0.9)
    tbl.columns[3].width = Inches(0.9)
    tbl.columns[4].width = Inches(0.8)
    
    table_data = [
        ["Feature / Capability", "SatQuery AI", "GEE (Google)", "QGIS / ArcGIS", "Standard DL"],
        ["Bi-Temporal Change Detection", "☑", "☑ (Script)", "☑ (Manual)", "☒"],
        ["Spectral Deltas (NDVI, NDWI, NDBI)", "☑", "☑", "☑ (Manual)", "☒"],
        ["Optical-SAR Multi-Modal Fusion", "☑", "☒ (Custom)", "☒", "☒"],
        ["Audit-Grade Evidence Extraction", "☑", "☒", "☒", "☒"],
        ["100% Explainable Trace Log", "☑", "☒", "☒", "☒"],
        ["Contradiction Penalty & Audit", "☑", "☒", "☒", "☒"],
        ["Zero-Hallucination RS-VQA", "☑", "☒", "☒", "☒"],
        ["Mission Control Web Dashboard", "☑", "☒ (Editor)", "☒ (Desktop)", "☒"]
    ]
    
    for r_idx, row in enumerate(table_data):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx, c_idx)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.alignment = PP_ALIGN.LEFT if c_idx == 0 else PP_ALIGN.CENTER
            
            # Header Row Styling
            if r_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = c_primary if c_idx != 1 else c_accent_blue
                p.font.bold = True
                p.font.size = Pt(8.5)
                p.font.color.rgb = c_white
            else:
                # Row styling
                cell.fill.solid()
                if c_idx == 1:
                    cell.fill.fore_color.rgb = RGBColor(240, 253, 244) # Highlight SatQuery column in light green
                else:
                    cell.fill.fore_color.rgb = RGBColor(255, 255, 255) if r_idx % 2 == 1 else RGBColor(248, 250, 252)
                
                p.font.size = Pt(8)
                if c_idx == 1:
                    p.font.bold = True
                    p.font.color.rgb = c_green
                elif val.startswith("☑"):
                    p.font.color.rgb = c_green
                elif val.startswith("☒"):
                    p.font.color.rgb = c_red
                else:
                    p.font.color.rgb = c_text_dark

    # Bottom Live Demo Bar
    demo_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(6.05), Inches(5.8), Inches(0.85)
    )
    demo_box.fill.solid()
    demo_box.fill.fore_color.rgb = RGBColor(238, 242, 255)
    demo_box.line.color.rgb = c_accent_blue
    demo_box.line.width = Pt(1.5)
    
    d_tf = demo_box.text_frame
    d_tf.word_wrap = True
    
    p_d1 = d_tf.paragraphs[0]
    p_d1.text = "💻 Live GitHub: https://github.com/navneetpatel9123-blip/SatQuery.git"
    p_d1.font.bold = True
    p_d1.font.size = Pt(9.5)
    p_d1.font.color.rgb = c_primary
    
    p_d2 = d_tf.add_paragraph()
    p_d2.text = "🚀 Interactive Dashboard Demo: http://localhost:3000/upload"
    p_d2.font.bold = True
    p_d2.font.size = Pt(9.5)
    p_d2.font.color.rgb = RGBColor(5, 150, 105)
    p_d2.space_before = Pt(2)
    
    prs.save(output_path)
    print(f"PowerPoint slide successfully created at: {output_path}")

if __name__ == "__main__":
    out_pptx = r"d:\satQUREY AI PROJECT\SatQuery_AI_SIH_References_and_Comparison_Slide.pptx"
    create_sih_slide(out_pptx)
